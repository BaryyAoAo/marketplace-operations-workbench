import base64
import io
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

import app as workbench


ONE_PIXEL_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVQIHWP4z8DwHwAFgAI/ScL/nwAAAABJRU5ErkJggg=="
)


def test_product_template_download_has_empty_import_sheet():
    response = TestClient(workbench.app).get("/api/products/template")
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content), read_only=True, data_only=True)
    assert book.sheetnames == ["商品导入", "填写示例"]
    rows = list(book.active.iter_rows(values_only=True))
    assert list(rows[0]) == workbench.FIELDS + workbench.OPTIONAL_PRODUCT_FIELDS
    assert not any(any(value is not None for value in row) for row in rows[1:])


def test_synthetic_sample_can_enter_publish_queue_after_image_import(tmp_path, monkeypatch):
    monkeypatch.setattr(workbench, "DATA", tmp_path)
    monkeypatch.setattr(workbench, "DB", tmp_path / "workbench.db")
    workbench.init_db()
    client = TestClient(workbench.app)
    sample = Path(workbench.ROOT / "sample_products_us.csv").read_text(encoding="utf-8")
    assert client.post("/api/products/import", json={"csv_text": sample}).json()["imported"] == 3
    product = next(p for p in client.get("/api/state").json()["products"] if p["sku"] == "US-RNG-001")
    uploaded = client.post(f"/api/products/{product['id']}/image", json={
        "filename": "US-RNG-001.png", "base64_data": base64.b64encode(ONE_PIXEL_PNG).decode()})
    assert uploaded.status_code == 200
    assert uploaded.json()["image_ready"] is True
    monkeypatch.setattr(workbench.gateway, "status", lambda: {"state": "online"})
    async def fake_run(_task_ids):
        return None
    monkeypatch.setattr(workbench, "_run_browser_publish_batch", fake_run)
    queued = client.post("/api/browser/publish/start", json={
        "product_ids": [product["id"]], "account_keys": ["chrome-default"], "confirm": True})
    assert queued.status_code == 200
    assert queued.json()["state"] == "queued"


def test_uncertain_task_can_be_verified_from_seller_list_without_item_url(tmp_path, monkeypatch):
    monkeypatch.setattr(workbench, "DATA", tmp_path)
    monkeypatch.setattr(workbench, "DB", tmp_path / "workbench.db")
    workbench.init_db()
    with workbench.connect() as con:
        con.execute("""INSERT INTO products(sku,name,category,condition,price_minor,currency,stock,
            material,size,color,image_url,source_note,created_at) VALUES
            ('T-1','Necklace','Jewellery','New',350000,'USD',1,'Unknown','One size','Gold','','sample','now')""")
        con.execute("""INSERT INTO drafts(product_id,title,description,language,status,version,created_at)
            VALUES(1,'Gold-tone pendant necklace — test item','Sample','en','approved',1,'now')""")
        con.execute("""INSERT INTO publish_tasks(draft_id,account_alias,status,created_at,updated_at)
            VALUES(1,'chrome-default','uncertain','now','now')""")
    async def fake_verify(_title):
        return {"state": "posted", "note": "Facebook 销售列表已显示此商品；平台仍在审核"}
    monkeypatch.setattr(workbench.gateway, "verify_listing", fake_verify)
    response = TestClient(workbench.app).post("/api/browser/publish/1/verify")
    assert response.status_code == 200
    assert response.json()["state"] == "posted"
    task = TestClient(workbench.app).get("/api/state").json()["publish_tasks"][0]
    assert task["status"] == "posted" and task["marketplace_url"] is None


def test_new_incoming_preview_creates_workbench_attention_flag(tmp_path, monkeypatch):
    monkeypatch.setattr(workbench, "DATA", tmp_path)
    monkeypatch.setattr(workbench, "DB", tmp_path / "workbench.db")
    workbench.init_db()
    client = TestClient(workbench.app)
    previews = ["You: Hello", "Buyer: Is it available?"]

    async def fake_refresh():
        if not previews:
            return []
        return [{"external_id": "123", "url": "https://www.facebook.com/messages/t/123/",
                 "title": "Buyer · Ring", "preview": previews.pop(0)}]

    monkeypatch.setattr(workbench.gateway, "refresh_conversations", fake_refresh)
    monkeypatch.setattr(workbench.gateway, "status", lambda: {"state": "online", "account_key": "chrome-default", "display_name": "Test Account"})
    first = client.post("/api/browser/conversations/refresh")
    assert first.json()["items"][0]["needs_attention"] == 0
    second = client.post("/api/browser/conversations/refresh")
    assert second.json()["items"][0]["needs_attention"] == 1
    conversation_id = second.json()["items"][0]["id"]
    assert client.post(f"/api/browser/conversations/{conversation_id}/handled").json() == {
        "needs_attention": False}
    previews.clear()
    empty = client.post("/api/browser/conversations/refresh").json()
    assert empty["visible_count"] == 0 and empty["items"] == []
    assert len(client.get("/api/browser/conversations").json()) == 1


def test_sheet_without_optional_image_url_column_can_be_imported(tmp_path, monkeypatch):
    monkeypatch.setattr(workbench, "DATA", tmp_path)
    monkeypatch.setattr(workbench, "DB", tmp_path / "workbench.db")
    workbench.init_db()
    csv_text = ("sku,name,category,condition,price,currency,stock,material,size,color,source_note,listing_title,listing_description\n"
                "US-TEST-001,Ring,Jewellery,New,25.00,USD,1,Unknown,US 7,Gold tone,staff sheet,Gold tone ring,Size US 7.\n")
    response = TestClient(workbench.app).post("/api/products/import", json={"csv_text": csv_text})
    assert response.status_code == 200
    assert response.json()["imported"] == 1


def test_interrupted_browser_batch_requires_manual_review(tmp_path, monkeypatch):
    monkeypatch.setattr(workbench, "DATA", tmp_path)
    monkeypatch.setattr(workbench, "DB", tmp_path / "workbench.db")
    workbench.init_db()
    with workbench.connect() as con:
        con.execute("""INSERT INTO products(sku,name,category,condition,price_minor,currency,stock,
            material,size,color,image_url,source_note,created_at) VALUES
            ('T-1','Ring','Jewellery','New',2500,'USD',1,'Unknown','US 7','Gold','','staff','now')""")
        con.execute("""INSERT INTO drafts(product_id,title,description,language,status,version,created_at)
            VALUES(1,'Ring','US 7','en','approved',1,'now')""")
        con.execute("""INSERT INTO publish_tasks(draft_id,account_alias,status,created_at,updated_at)
            VALUES(1,'chrome-default','preparing','now','now')""")
        workbench.event(con, "browser_batch_queued", "publish_task", 1)
    workbench.recover_interrupted_browser_batches()
    with workbench.connect() as con:
        row = con.execute("SELECT status,note FROM publish_tasks WHERE id=1").fetchone()
    assert row["status"] == "uncertain"
    assert "人工核对" in row["note"]


def test_uncertain_browser_send_is_not_retried_with_same_text(tmp_path, monkeypatch):
    monkeypatch.setattr(workbench, "DATA", tmp_path)
    monkeypatch.setattr(workbench, "DB", tmp_path / "workbench.db")
    workbench.init_db()
    with workbench.connect() as con:
        con.execute("""INSERT INTO browser_conversations
            (account_key,external_id,url,title,first_seen_at,last_seen_at)
            VALUES('chrome-default','123','https://www.facebook.com/messages/t/123/',
                   'Seller · Ring','now','now')""")
    monkeypatch.setattr(workbench.gateway, "status", lambda: {"state": "online", "account_key": "chrome-default"})
    calls = []

    async def uncertain_send(url, message):
        calls.append((url, message))
        return {"state": "uncertain"}

    monkeypatch.setattr(workbench.gateway, "send_message", uncertain_send)
    client = TestClient(workbench.app)
    first = client.post("/api/browser/conversations/1/send", json={"message": "Is pickup possible?"})
    second = client.post("/api/browser/conversations/1/send", json={"message": "Is pickup possible?"})
    assert first.status_code == 200 and first.json()["state"] == "uncertain"
    assert second.status_code == 409
    assert len(calls) == 1


def test_ai_reply_uses_entire_read_history_and_requires_latest_incoming(tmp_path, monkeypatch):
    monkeypatch.setattr(workbench, "DATA", tmp_path)
    monkeypatch.setattr(workbench, "DB", tmp_path / "workbench.db")
    workbench.init_db()
    with workbench.connect() as con:
        con.execute("""INSERT INTO products(sku,name,category,condition,price_minor,currency,stock,
            material,size,color,image_url,source_note,created_at,listing_title,listing_description)
            VALUES('R-1','Ring','ring','new',2500,'USD',1,'steel','US 7','silver','','staff','now',
                   'Silver ring','Steel band, size US 7.')""")
        con.execute("""INSERT INTO browser_conversations
            (account_key,external_id,url,title,first_seen_at,last_seen_at)
            VALUES('chrome-default','123','https://www.facebook.com/messages/t/123/',
                   'Buyer · Ring','now','now')""")
        for direction, sender, text in [
            ('in','Buyer','Is it available?'), ('out','You','Yes.'),
            ('in','Buyer','Can you ship it?'),
        ]:
            con.execute("""INSERT INTO browser_messages
                (conversation_id,fingerprint,time_label,sender,direction,text,fetched_at)
                VALUES(1,?,?,?,?,?,'now')""", (text, '10:00', sender, direction, text))
    monkeypatch.setattr(workbench.gateway, "status", lambda: {"state":"online", "account_key":"chrome-default"})
    monkeypatch.setattr(workbench.ai, "plan_reply", lambda question, title, recent: (
        workbench.ai.ReplyPlan(intent="shipping", use_product=True, use_history=True,
                               policy_topics=["shipping"], product_query="Ring"), "test-model"))
    captured = {}
    def fake_reply(product, question, *, conversation_history, knowledge_context):
        captured["facts"] = workbench.ai.product_facts(product)
        captured["question"] = question
        captured["history"] = conversation_history
        captured["knowledge"] = knowledge_context
        return workbench.ai.InquiryReply(intent="delivery", answer="I will confirm shipping options.", review_notes=[]), "test-model"
    monkeypatch.setattr(workbench.ai, "inquiry_reply", fake_reply)
    client = TestClient(workbench.app)
    result = client.post('/api/browser/conversations/1/reply-ai', json={
        'allow_external_processing':True})
    assert result.status_code == 200
    assert captured["question"] == "Can you ship it?"
    assert [message["text"] for message in captured["history"]] == [
        "Is it available?", "Yes.", "Can you ship it?"]
    assert captured["facts"]["listing_description"] == "Steel band, size US 7."
    assert captured["knowledge"]["policies"]["shipping"] == workbench.ai_knowledge.DEFAULT_POLICIES["shipping"]
    assert "配送规则" in result.json()["basis"]["policies"]
    assert result.json()["basis"]["history_count"] == 3
    with workbench.connect() as con:
        con.execute("""INSERT INTO browser_messages
            (conversation_id,fingerprint,time_label,sender,direction,text,fetched_at)
            VALUES(1,'last','10:03','You','out','I will check.','now')""")
    blocked = client.post('/api/browser/conversations/1/reply-ai', json={
        'product_id':1, 'allow_external_processing':True})
    assert blocked.status_code == 409


def test_conversation_history_is_not_read_from_another_browser_account(tmp_path, monkeypatch):
    monkeypatch.setattr(workbench, "DATA", tmp_path)
    monkeypatch.setattr(workbench, "DB", tmp_path / "workbench.db")
    workbench.init_db()
    with workbench.connect() as con:
        con.execute("""INSERT INTO browser_conversations
            (account_key,external_id,url,title,first_seen_at,last_seen_at)
            VALUES('other-profile','999','https://www.facebook.com/messages/t/999/',
                   'Other account conversation','now','now')""")
    monkeypatch.setattr(workbench.gateway, "status", lambda: {"state": "online", "account_key": "chrome-default"})
    client = TestClient(workbench.app)
    assert client.get("/api/browser/conversations").json()[0]["account_key"] == "other-profile"
    assert client.post("/api/browser/conversations/1/read").status_code == 409
    assert client.post("/api/browser/conversations/1/send", json={"message": "Hello"}).status_code == 409
