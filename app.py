"""Local Marketplace workbench with an optional, operator controlled Chrome adapter."""

from __future__ import annotations

import asyncio
import csv
import base64
import binascii
import hashlib
import io
import json
import os
import re
import sqlite3
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Literal
from urllib.parse import urlparse

from fastapi import FastAPI, HTTPException
from fastapi.responses import FileResponse
from openpyxl import load_workbook
from openai import OpenAIError
from pydantic import BaseModel, Field

import ai
import ai_knowledge
import ai_settings
import reply_agent
from browser_gateway import gateway


ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
DB = DATA / "workbench.db"
PRODUCT_TEMPLATE = ROOT / "outputs" / "product-template" / "商品导入模板.xlsx"
FIELDS = ["sku", "name", "category", "condition", "price", "currency", "stock", "material", "size", "color", "image_url", "source_note"]
REQUIRED_PRODUCT_FIELDS = [field for field in FIELDS if field != "image_url"]
OPTIONAL_PRODUCT_FIELDS = ["listing_title", "listing_description"]
UNSUPPORTED_CLAIMS = ("solid gold", "18k", "24k", "diamond", "hypoallergenic", "waterproof", "sterling silver", "handmade", "authentic", "free shipping", "guaranteed delivery")
app = FastAPI(title="Marketplace Workbench", version="0.3.42")


def now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def connect() -> sqlite3.Connection:
    DATA.mkdir(exist_ok=True)
    con = sqlite3.connect(DB)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA foreign_keys = ON")
    return con


def init_db() -> None:
    with connect() as con:
        con.executescript("""
        CREATE TABLE IF NOT EXISTS products (
          id INTEGER PRIMARY KEY, sku TEXT NOT NULL, name TEXT NOT NULL,
          category TEXT NOT NULL, condition TEXT NOT NULL, price_minor INTEGER NOT NULL,
          currency TEXT NOT NULL, stock INTEGER NOT NULL, material TEXT NOT NULL,
          size TEXT NOT NULL, color TEXT NOT NULL, image_url TEXT NOT NULL,
          source_note TEXT NOT NULL, created_at TEXT NOT NULL,
          active INTEGER NOT NULL DEFAULT 1
        );
        CREATE TABLE IF NOT EXISTS drafts (
          id INTEGER PRIMARY KEY, product_id INTEGER NOT NULL REFERENCES products(id),
          title TEXT NOT NULL, description TEXT NOT NULL, language TEXT NOT NULL,
          status TEXT NOT NULL, version INTEGER NOT NULL, created_at TEXT NOT NULL,
          approved_at TEXT
        );
        CREATE TABLE IF NOT EXISTS publish_tasks (
          id INTEGER PRIMARY KEY, draft_id INTEGER NOT NULL REFERENCES drafts(id),
          account_alias TEXT NOT NULL, status TEXT NOT NULL, marketplace_url TEXT,
          note TEXT NOT NULL DEFAULT '', created_at TEXT NOT NULL, updated_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS inquiries (
          id INTEGER PRIMARY KEY, task_id INTEGER NOT NULL REFERENCES publish_tasks(id),
          buyer_alias TEXT NOT NULL, question TEXT NOT NULL, reply_draft TEXT NOT NULL,
          reply_status TEXT NOT NULL, stage TEXT NOT NULL, channel TEXT NOT NULL,
          next_followup_at TEXT, created_at TEXT NOT NULL, updated_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS audit_events (
          id INTEGER PRIMARY KEY, event_type TEXT NOT NULL, entity_type TEXT NOT NULL,
          entity_id INTEGER NOT NULL, detail TEXT NOT NULL, created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS browser_conversations (
          id INTEGER PRIMARY KEY, account_key TEXT NOT NULL, external_id TEXT NOT NULL,
          url TEXT NOT NULL, title TEXT NOT NULL, last_preview TEXT NOT NULL DEFAULT '',
          needs_attention INTEGER NOT NULL DEFAULT 0, first_seen_at TEXT NOT NULL,
          last_seen_at TEXT NOT NULL, UNIQUE(account_key, external_id)
        );
        CREATE TABLE IF NOT EXISTS browser_messages (
          id INTEGER PRIMARY KEY, conversation_id INTEGER NOT NULL REFERENCES browser_conversations(id),
          fingerprint TEXT NOT NULL, time_label TEXT NOT NULL, sender TEXT NOT NULL,
          direction TEXT NOT NULL, text TEXT NOT NULL, fetched_at TEXT NOT NULL,
          UNIQUE(conversation_id, fingerprint)
        );
        CREATE TABLE IF NOT EXISTS browser_send_attempts (
          id INTEGER PRIMARY KEY, conversation_id INTEGER NOT NULL REFERENCES browser_conversations(id),
          body TEXT NOT NULL, state TEXT NOT NULL, created_at TEXT NOT NULL, updated_at TEXT NOT NULL
        );
        """)
        for table, columns in {
            "products": {"listing_title": "TEXT NOT NULL DEFAULT ''", "listing_description": "TEXT NOT NULL DEFAULT ''",
                         "image_path": "TEXT NOT NULL DEFAULT ''", "image_hash": "TEXT NOT NULL DEFAULT ''"},
            "drafts": {"generation_source": "TEXT NOT NULL DEFAULT 'template'", "model_id": "TEXT", "review_notes": "TEXT NOT NULL DEFAULT '[]'"},
            "inquiries": {"reply_source": "TEXT NOT NULL DEFAULT 'template'", "reply_model": "TEXT", "ai_notes": "TEXT NOT NULL DEFAULT '[]'"},
        }.items():
            present = {row[1] for row in con.execute(f"PRAGMA table_info({table})")}
            for column, definition in columns.items():
                if column not in present:
                    con.execute(f"ALTER TABLE {table} ADD COLUMN {column} {definition}")
        con.commit()
        # Older databases have a table-level UNIQUE(sku). Rebuild once so an archived
        # product can retain its history while a new import reuses that SKU.
        product_sql = con.execute("SELECT sql FROM sqlite_master WHERE type='table' AND name='products'").fetchone()[0]
        if "sku TEXT UNIQUE" in product_sql:
            con.execute("PRAGMA foreign_keys = OFF")
            try:
                con.execute("BEGIN IMMEDIATE")
                con.execute("""CREATE TABLE products_new (
                  id INTEGER PRIMARY KEY, sku TEXT NOT NULL, name TEXT NOT NULL,
                  category TEXT NOT NULL, condition TEXT NOT NULL, price_minor INTEGER NOT NULL,
                  currency TEXT NOT NULL, stock INTEGER NOT NULL, material TEXT NOT NULL,
                  size TEXT NOT NULL, color TEXT NOT NULL, image_url TEXT NOT NULL,
                  source_note TEXT NOT NULL, created_at TEXT NOT NULL,
                  listing_title TEXT NOT NULL DEFAULT '', listing_description TEXT NOT NULL DEFAULT '',
                  image_path TEXT NOT NULL DEFAULT '', image_hash TEXT NOT NULL DEFAULT '',
                  active INTEGER NOT NULL DEFAULT 1
                )""")
                columns = ("id,sku,name,category,condition,price_minor,currency,stock,material,size,color,"
                           "image_url,source_note,created_at,listing_title,listing_description,image_path,image_hash,active")
                con.execute(f"INSERT INTO products_new ({columns}) SELECT {columns.replace(',active', ',1')} FROM products")
                con.execute("DROP TABLE products")
                con.execute("ALTER TABLE products_new RENAME TO products")
                con.commit()
            except Exception:
                con.rollback()
                raise
            finally:
                con.execute("PRAGMA foreign_keys = ON")
            if con.execute("PRAGMA foreign_key_check").fetchone():
                raise RuntimeError("商品表迁移后外键校验失败")
        con.execute("CREATE UNIQUE INDEX IF NOT EXISTS products_active_sku ON products(sku) WHERE active=1")


init_db()


def recover_interrupted_browser_batches() -> None:
    """A process restart must never silently repeat a listing submission."""
    with connect() as con:
        rows = con.execute("""SELECT t.id,t.status FROM publish_tasks t
            WHERE t.status IN ('queued','preparing') AND EXISTS (
              SELECT 1 FROM audit_events e WHERE e.entity_type='publish_task'
              AND e.entity_id=t.id AND e.event_type='browser_batch_queued')""").fetchall()
        for row in rows:
            status = 'uncertain' if row['status'] == 'preparing' else 'needs_attention'
            note = ('服务重启时网页发布进行中；可能已提交，请人工核对后处理'
                    if row['status'] == 'preparing' else '服务重启后队列暂停；请人工检查后重新安排')
            con.execute("UPDATE publish_tasks SET status=?,note=?,updated_at=? WHERE id=?",
                        (status, note, now(), row['id']))
            con.execute("""INSERT INTO audit_events(event_type,entity_type,entity_id,detail,created_at)
                VALUES(?,?,?,?,?)""", ('browser_batch_recovered', 'publish_task', row['id'],
                                         status, now()))


recover_interrupted_browser_batches()


def event(con: sqlite3.Connection, event_type: str, entity_type: str, entity_id: int, detail: str = "") -> None:
    con.execute("INSERT INTO audit_events(event_type,entity_type,entity_id,detail,created_at) VALUES(?,?,?,?,?)",
                (event_type, entity_type, entity_id, detail, now()))


def one(con: sqlite3.Connection, table: str, item_id: int) -> sqlite3.Row:
    if table not in {"products", "drafts", "publish_tasks", "inquiries"}:
        raise ValueError("Invalid table")
    row = con.execute(f"SELECT * FROM {table} WHERE id=?", (item_id,)).fetchone()
    if row is None:
        raise HTTPException(404, f"{table} item not found")
    return row


def clean(value: str | None) -> str:
    return (value or "").strip()


def validate_row(raw: dict[str, str]) -> tuple[dict, list[str]]:
    values = {key: clean(raw.get(key)) for key in FIELDS + OPTIONAL_PRODUCT_FIELDS}
    errors = [f"{key} 为空" for key in REQUIRED_PRODUCT_FIELDS if not values[key]]
    price_minor = 0
    stock = 0
    try:
        price = Decimal(values["price"])
        if not price.is_finite() or price <= 0 or price.as_tuple().exponent < -2:
            raise InvalidOperation
        price_minor = int(price * 100)
    except (InvalidOperation, ValueError):
        errors.append("price 必须是大于 0、最多两位小数的金额")
    try:
        stock = int(values["stock"])
        if stock < 0:
            raise ValueError
    except ValueError:
        errors.append("stock 必须是非负整数")
    if values["currency"].upper() != "USD":
        errors.append("当前洛杉矶演示只接受 USD")
    if values["image_url"] and urlparse(values["image_url"]).scheme not in {"http", "https"}:
        errors.append("image_url 必须是 http(s) 地址")
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_-]{1,63}", values["sku"]):
        errors.append("sku 只允许字母、数字、下划线、连字符，长度 2–64")
    values["price_minor"] = price_minor
    values["stock"] = stock
    values["currency"] = values["currency"].upper()
    return values, errors


def import_records(con: sqlite3.Connection, rows: list[tuple[int, dict]]) -> dict:
    result: dict = {"imported": 0, "errors": []}
    seen: set[str] = set()
    for line_no, raw in rows:
        values, errors = validate_row(raw)
        if values["sku"] in seen:
            errors.append("本次文件内 SKU 重复")
        seen.add(values["sku"])
        if con.execute("SELECT 1 FROM products WHERE sku=? AND active=1", (values["sku"],)).fetchone():
            errors.append("SKU 已存在，已跳过")
        if errors:
            result["errors"].append({"line": line_no, "sku": values["sku"], "messages": errors})
            continue
        cur = con.execute("""INSERT INTO products
          (sku,name,category,condition,price_minor,currency,stock,material,size,color,image_url,source_note,
           listing_title,listing_description,created_at)
          VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
          tuple(values[key] for key in ["sku", "name", "category", "condition", "price_minor", "currency", "stock", "material", "size", "color", "image_url", "source_note", "listing_title", "listing_description"]) + (now(),))
        event(con, "product_imported", "product", cur.lastrowid, values["sku"])
        result["imported"] += 1
    return result


def import_csv(con: sqlite3.Connection, csv_text: str) -> dict:
    try:
        reader = csv.DictReader(io.StringIO(csv_text.lstrip("\ufeff")))
        if not reader.fieldnames or any(field not in reader.fieldnames for field in REQUIRED_PRODUCT_FIELDS):
            raise HTTPException(400, "CSV 缺少必填列：" + ", ".join(REQUIRED_PRODUCT_FIELDS))
        return import_records(con, list(enumerate(reader, start=2)))
    except csv.Error as exc:
        raise HTTPException(400, f"CSV 无法解析：{exc}") from exc


class ImportRequest(BaseModel):
    csv_text: str = Field(min_length=1)


class XlsxImportRequest(BaseModel):
    base64_data: str = Field(min_length=1)


class ClearProductsRequest(BaseModel):
    confirm: Literal[True]


class DraftEdit(BaseModel):
    title: str = Field(min_length=1, max_length=200)
    description: str = Field(min_length=1, max_length=5000)


class QueueRequest(BaseModel):
    account_alias: str = Field(min_length=2, max_length=50, pattern=r"^[A-Za-z0-9_-]+$")


class TaskResult(BaseModel):
    status: Literal["posted", "failed", "needs_attention"]
    marketplace_url: str = ""
    note: str = ""


class InquiryCreate(BaseModel):
    task_id: int
    buyer_alias: str = Field(min_length=1, max_length=80)
    question: str = Field(min_length=1, max_length=1000)


class ReplyEdit(BaseModel):
    reply_draft: str = Field(min_length=1, max_length=2000)
    reply_status: Literal["draft", "approved", "sent_manually"]


class FollowupEdit(BaseModel):
    stage: Literal["new", "responded", "negotiating", "won", "lost"]
    channel: Literal["marketplace", "whatsapp_manual", "other_manual"]
    next_followup_at: str | None = None


class AIReplyRequest(BaseModel):
    allow_external_processing: bool = False


class AISettingsUpdate(BaseModel):
    provider: Literal["siliconflow", "compatible", "openai"]
    base_url: str = Field(default="", max_length=500)
    model: str = Field(min_length=1, max_length=200)
    api_key: str = Field(default="", max_length=1000)
    reply_system_prompt: str = Field(default="", max_length=6000)
    clear_api_key: bool = False


class QuickReplyInput(BaseModel):
    intent: Literal["general", "availability", "price", "size", "shipping", "pickup", "payment", "returns", "negotiation"]
    title: str = Field(min_length=1, max_length=80)
    content: str = Field(min_length=1, max_length=1200)


class AIKnowledgeUpdate(BaseModel):
    policies: dict[str, str]
    quick_replies: list[QuickReplyInput] = Field(max_length=20)


class BrowserSendRequest(BaseModel):
    message: str = Field(min_length=1, max_length=2000)


class BrowserAIReplyRequest(BaseModel):
    product_id: int | None = None
    allow_external_processing: bool = False


class ProductImageUpload(BaseModel):
    filename: str = Field(min_length=1, max_length=255)
    base64_data: str = Field(min_length=1, max_length=20_000_000)


class PublishBatchRequest(BaseModel):
    product_ids: list[int] = Field(min_length=1, max_length=100)
    account_keys: list[str] = Field(min_length=1, max_length=10)
    confirm: bool = False


def draft_issues(product: sqlite3.Row, title: str, description: str) -> list[str]:
    issues = []
    if len(title) > 100:
        issues.append("标题超过 100 字符")
    if len(description) > 1000:
        issues.append("描述超过 1000 字符")
    facts = " ".join(str(product[key]) for key in ("material", "size", "color")).lower()
    copy = f"{title} {description}".lower()
    for claim in UNSUPPORTED_CLAIMS:
        if claim in copy and claim not in facts:
            issues.append(f"卖点“{claim}”没有商品字段依据")
    currency = product["currency"].upper()
    currency_codes = {code.upper() for code in re.findall(r"\b(?:USD|QAR|EUR|GBP|SAR|AED)\b", copy, re.I)}
    if currency_codes - {currency}:
        issues.append(f"文案出现与 {currency} 不一致的币种")
    amount_pattern = rf"(?:{re.escape(currency)}\s*\d+(?:\.\d{{1,2}})?|\d+(?:\.\d{{1,2}})?\s*{re.escape(currency)})"
    for match in re.finditer(amount_pattern, copy, re.I):
        amount = re.search(r"\d+(?:\.\d{1,2})?", match.group()).group()
        if int(Decimal(amount) * 100) != product["price_minor"]:
            issues.append("文案价格与结构化价格不一致")
    if currency == "USD":
        for match in re.finditer(r"\$\s*(\d+(?:\.\d{1,2})?)", copy):
            if int(Decimal(match.group(1)) * 100) != product["price_minor"]:
                issues.append("文案价格与结构化价格不一致")
    for match in re.finditer(r"\bprice\s*(?:is|:)?\s*(\d+(?:\.\d{1,2})?)\b", copy, re.I):
        if int(Decimal(match.group(1)) * 100) != product["price_minor"]:
            issues.append("文案价格与结构化价格不一致")
    return issues


def reply_issues(product: sqlite3.Row, question: str, answer: str) -> list[str]:
    issues = draft_issues(product, "", answer)
    if any(word in question.lower() for word in ("price", "cost", "how much")):
        expected = f"{Decimal(product['price_minor']) / 100:.2f} {product['currency']}"
        if expected.lower() not in answer.lower():
            issues.append("询问价格时必须使用结构化价格")
    if product["stock"] == 0 and re.search(r"\b(in stock|available now|is (?:currently |still )?available)\b", answer, re.I):
        issues.append("回复库存状态与结构化库存不一致")
    if product["stock"] > 0 and re.search(r"\b(?:out of stock|currently unavailable)\b", answer, re.I):
        issues.append("回复库存状态与结构化库存不一致")
    return issues


@app.get("/")
def index():
    return FileResponse(ROOT / "static" / "index.html")


@app.get("/inbox")
def inbox_page():
    return FileResponse(ROOT / "static" / "index.html")


@app.get("/ai-settings")
def ai_settings_page():
    return FileResponse(ROOT / "static" / "index.html")


@app.get("/styles.css")
def styles():
    return FileResponse(ROOT / "static" / "styles.css")


@app.get("/main.js")
def script():
    return FileResponse(ROOT / "static" / "main.js")


@app.get("/api/state")
def state():
    with connect() as con:
        result = {table: [dict(row) for row in con.execute(f"SELECT * FROM {table} ORDER BY id DESC")]
                  for table in ("drafts", "publish_tasks", "inquiries", "audit_events")}
        result["products"] = [dict(row) for row in con.execute("SELECT * FROM products WHERE active=1 ORDER BY id DESC")]
        result["archived_products"] = [dict(row) for row in con.execute("SELECT * FROM products WHERE active=0 ORDER BY id DESC")]
        return result


@app.get("/api/ai/status")
def ai_status():
    return ai.public_status()


@app.put("/api/ai/settings")
def update_ai_settings(payload: AISettingsUpdate):
    try:
        ai_settings.save(**payload.model_dump())
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc
    return ai.public_status()


@app.delete("/api/ai/settings")
def reset_ai_settings():
    ai_settings.clear()
    return ai.public_status()


@app.get("/api/ai/knowledge")
def get_ai_knowledge():
    return ai_knowledge.load()


@app.put("/api/ai/knowledge")
def update_ai_knowledge(payload: AIKnowledgeUpdate):
    try:
        return ai_knowledge.save(payload.policies, [item.model_dump() for item in payload.quick_replies])
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc


@app.get("/api/products/template")
def products_template():
    return FileResponse(
        PRODUCT_TEMPLATE,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="商品导入模板.xlsx",
    )


@app.post("/api/products/import")
def products_import(payload: ImportRequest):
    with connect() as con:
        return import_csv(con, payload.csv_text)


@app.post("/api/products/clear")
def products_clear(payload: ClearProductsRequest):
    with connect() as con:
        busy = con.execute("SELECT 1 FROM publish_tasks WHERE status IN ('queued','preparing') LIMIT 1").fetchone()
        if busy or (_batch_task is not None and not _batch_task.done()):
            raise HTTPException(409, "发布队列正在执行，结束后才能清空当前商品表")
        count = con.execute("SELECT COUNT(*) FROM products WHERE active=1").fetchone()[0]
        con.execute("UPDATE products SET active=0 WHERE active=1")
        event(con, "product_catalog_cleared", "catalog", 0, f"archived={count}")
        return {"cleared": count}


@app.post("/api/products/import-xlsx")
def products_import_xlsx(payload: XlsxImportRequest):
    try:
        raw = base64.b64decode(payload.base64_data, validate=True)
        if len(raw) > 5_000_000:
            raise HTTPException(413, "Excel 文件最多 5 MB")
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = book.active
        iterator = sheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(iterator)]
        if any(field not in headers for field in REQUIRED_PRODUCT_FIELDS):
            raise HTTPException(400, "Excel 首行缺少必填列：" + ", ".join(REQUIRED_PRODUCT_FIELDS))
        rows = [(line, {headers[index]: "" if value is None else str(value) for index, value in enumerate(cells) if index < len(headers)})
                for line, cells in enumerate(iterator, start=2) if any(value is not None for value in cells)]
    except HTTPException:
        raise
    except (ValueError, OSError, StopIteration) as exc:
        raise HTTPException(400, "无法读取 Excel 文件，请检查 .xlsx 格式") from exc
    with connect() as con:
        return import_records(con, rows)


@app.post("/api/products/{product_id}/image")
def product_image_upload(product_id: int, payload: ProductImageUpload):
    try:
        image_bytes = base64.b64decode(payload.base64_data, validate=True)
    except (ValueError, binascii.Error) as exc:
        raise HTTPException(422, "图片不是有效的 Base64 文件") from exc
    if not 0 < len(image_bytes) <= 12_000_000:
        raise HTTPException(422, "图片大小必须在 12 MB 以内")
    if image_bytes.startswith(b"\xff\xd8\xff"):
        extension = ".jpg"
    elif image_bytes.startswith(b"\x89PNG\r\n\x1a\n"):
        extension = ".png"
    elif image_bytes.startswith(b"RIFF") and image_bytes[8:12] == b"WEBP":
        extension = ".webp"
    else:
        raise HTTPException(422, "仅支持 JPEG、PNG 或 WebP 图片")
    with connect() as con:
        product = one(con, "products", product_id)
        digest = hashlib.sha256(image_bytes).hexdigest()
        image_dir = DATA / "product_images" / product["sku"]
        image_dir.mkdir(parents=True, exist_ok=True)
        image_path = image_dir / (digest[:20] + extension)
        if not image_path.exists():
            image_path.write_bytes(image_bytes)
        con.execute("UPDATE products SET image_path=?,image_hash=? WHERE id=?",
                    (str(image_path), digest, product_id))
        event(con, "product_image_uploaded", "product", product_id, f"sha256={digest[:12]}")
    return {"product_id": product_id, "image_ready": True, "sha256": digest}


@app.post("/api/products/{product_id}/draft")
def create_draft(product_id: int):
    with connect() as con:
        product = one(con, "products", product_id)
        version = con.execute("SELECT COALESCE(MAX(version),0)+1 FROM drafts WHERE product_id=?", (product_id,)).fetchone()[0]
        title = product["listing_title"] or f"{product['name']} · {product['color']}"
        description = product["listing_description"] or (
            f"{product['name']}. Material: {product['material']}. Size: {product['size']}. "
            f"Color: {product['color']}. Condition: {product['condition']}. "
            "Please confirm availability and delivery details before purchase.")
        cur = con.execute("INSERT INTO drafts(product_id,title,description,language,status,version,created_at) VALUES(?,?,?,?,?,?,?)",
                          (product_id, title, description, "en", "pending", version, now()))
        event(con, "draft_created", "draft", cur.lastrowid, f"product={product_id}, version={version}")
        return {"id": cur.lastrowid, "issues": draft_issues(product, title, description)}


@app.post("/api/products/{product_id}/draft-ai")
def create_ai_draft(product_id: int):
    with connect() as con:
        product = one(con, "products", product_id)
        try:
            copy, model = ai.listing_copy(product)
        except ai.AIUnavailable as exc:
            raise HTTPException(503, str(exc)) from exc
        except OpenAIError as exc:
            raise HTTPException(503, "模型服务调用失败，请检查本机配置与网络") from exc
        title = copy.title.strip()
        description = copy.description.strip()
        if not title or not description or len(title) > 100 or len(description) > 1000:
            raise HTTPException(422, "模型草稿为空或超过长度限制")
        issues = draft_issues(product, title, description)
        version = con.execute("SELECT COALESCE(MAX(version),0)+1 FROM drafts WHERE product_id=?", (product_id,)).fetchone()[0]
        cur = con.execute("""INSERT INTO drafts
          (product_id,title,description,language,status,version,created_at,generation_source,model_id,review_notes)
          VALUES(?,?,?,?,?,?,?,?,?,?)""",
          (product_id, title, description, "en", "pending", version, now(), "llm", model,
           json.dumps(copy.review_notes, ensure_ascii=False)))
        event(con, "ai_draft_created", "draft", cur.lastrowid, f"model={model}, product={product_id}")
        return {"id": cur.lastrowid, "issues": issues, "review_notes": copy.review_notes, "model": model}


@app.put("/api/drafts/{draft_id}")
def edit_draft(draft_id: int, payload: DraftEdit):
    with connect() as con:
        draft = one(con, "drafts", draft_id)
        if draft["status"] != "pending":
            raise HTTPException(409, "已批准草稿不可修改；请为商品建立新版本")
        con.execute("UPDATE drafts SET title=?,description=? WHERE id=?", (payload.title.strip(), payload.description.strip(), draft_id))
        event(con, "draft_edited", "draft", draft_id)
        return {"issues": draft_issues(one(con, "products", draft["product_id"]), payload.title, payload.description)}


@app.post("/api/drafts/{draft_id}/approve")
def approve_draft(draft_id: int):
    with connect() as con:
        draft = one(con, "drafts", draft_id)
        if draft["status"] != "pending":
            raise HTTPException(409, "草稿已批准")
        product = one(con, "products", draft["product_id"])
        issues = draft_issues(product, draft["title"], draft["description"])
        if issues:
            raise HTTPException(422, {"issues": issues})
        con.execute("UPDATE drafts SET status='approved',approved_at=? WHERE id=?", (now(), draft_id))
        event(con, "draft_approved", "draft", draft_id, "human confirmation")
        return {"status": "approved"}


@app.post("/api/drafts/{draft_id}/queue")
def queue_draft(draft_id: int, payload: QueueRequest):
    with connect() as con:
        draft = one(con, "drafts", draft_id)
        if draft["status"] != "approved":
            raise HTTPException(409, "只有人工批准的草稿可进入发布队列")
        cur = con.execute("INSERT INTO publish_tasks(draft_id,account_alias,status,created_at,updated_at) VALUES(?,?,?,?,?)",
                          (draft_id, payload.account_alias, "queued", now(), now()))
        event(con, "task_queued", "publish_task", cur.lastrowid, payload.account_alias)
        return {"id": cur.lastrowid}


@app.post("/api/tasks/{task_id}/result")
def task_result(task_id: int, payload: TaskResult):
    with connect() as con:
        task = one(con, "publish_tasks", task_id)
        if task["status"] == "posted":
            raise HTTPException(409, "已发布任务不可重复回填")
        url = payload.marketplace_url.strip()
        if payload.status == "posted":
            if not accepted_listing_url(url):
                raise HTTPException(422, "请填写 Facebook Marketplace 商品链接或 Facebook 商品分享链接")
        con.execute("UPDATE publish_tasks SET status=?,marketplace_url=?,note=?,updated_at=? WHERE id=?",
                    (payload.status, url if payload.status == "posted" else None, payload.note.strip(), now(), task_id))
        event(con, "publish_result_recorded", "publish_task", task_id, payload.status)
        return {"status": payload.status}


def reply_for(product: sqlite3.Row, question: str) -> str:
    q = question.lower()
    price = f"{Decimal(product['price_minor']) / 100:.2f} {product['currency']}"
    if any(word in q for word in ("price", "cost", "how much")):
        detail = f"The listed price is {price}."
    elif any(word in q for word in ("size", "fit")):
        detail = f"The listed size is {product['size']}."
    elif any(word in q for word in ("material", "metal", "made of")):
        detail = f"The listed material is {product['material']}."
    elif any(word in q for word in ("available", "stock")):
        detail = "It is currently marked in stock." if product["stock"] > 0 else "It is currently marked out of stock."
    else:
        detail = f"The listing shows {product['material']}, size {product['size']}, color {product['color']}."
    return f"Hi, thanks for asking about {product['name']}. {detail} Please confirm current availability and delivery details before ordering."


def accepted_listing_url(url: str) -> bool:
    """Accept canonical Marketplace item URLs and Facebook's own share links.

    A share link is an operator-provided record; the app does not assert that
    every possible share target is a Marketplace item.
    """
    parsed = urlparse(url)
    if parsed.scheme != "https" or parsed.hostname not in {"facebook.com", "www.facebook.com", "m.facebook.com"}:
        return False
    if parsed.username or parsed.password:
        return False
    return bool(re.fullmatch(r"/marketplace/item/\d+/?", parsed.path) or
                re.fullmatch(r"/share/[A-Za-z0-9]+/?", parsed.path))


@app.post("/api/inquiries")
def create_inquiry(payload: InquiryCreate):
    with connect() as con:
        task = one(con, "publish_tasks", payload.task_id)
        draft = one(con, "drafts", task["draft_id"])
        product = one(con, "products", draft["product_id"])
        cur = con.execute("""INSERT INTO inquiries
          (task_id,buyer_alias,question,reply_draft,reply_status,stage,channel,created_at,updated_at)
          VALUES(?,?,?,?,?,?,?,?,?)""",
          (payload.task_id, payload.buyer_alias.strip(), payload.question.strip(), reply_for(product, payload.question),
           "draft", "new", "marketplace", now(), now()))
        event(con, "inquiry_recorded", "inquiry", cur.lastrowid, f"task={payload.task_id}")
        return {"id": cur.lastrowid}


@app.post("/api/inquiries/{inquiry_id}/reply-ai")
def create_ai_reply(inquiry_id: int, payload: AIReplyRequest):
    if not payload.allow_external_processing:
        raise HTTPException(422, "请先确认询盘已脱敏并允许发送给配置的模型服务")
    with connect() as con:
        inquiry = one(con, "inquiries", inquiry_id)
        if inquiry["reply_status"] == "sent_manually":
            raise HTTPException(409, "已记录为发送的询盘不可重新生成草稿")
        task = one(con, "publish_tasks", inquiry["task_id"])
        draft = one(con, "drafts", task["draft_id"])
        product = one(con, "products", draft["product_id"])
        knowledge = ai_knowledge.select(inquiry["question"])
        try:
            reply, model = ai.inquiry_reply(product, inquiry["question"], knowledge_context=knowledge)
        except ai.AIUnavailable as exc:
            raise HTTPException(503, str(exc)) from exc
        except OpenAIError as exc:
            raise HTTPException(503, "模型服务调用失败，请检查本机配置与网络") from exc
        answer = reply.answer.strip()
        if not answer or len(answer) > 2000:
            raise HTTPException(422, "模型回复为空或超过长度限制")
        issues = reply_issues(product, inquiry["question"], answer)
        if issues:
            raise HTTPException(422, {"issues": issues})
        review_notes = list(dict.fromkeys([*reply.review_notes, *knowledge["missing_facts"]]))
        con.execute("""UPDATE inquiries SET reply_draft=?,reply_status='draft',reply_source='llm',
          reply_model=?,ai_notes=?,updated_at=? WHERE id=?""",
          (answer, model, json.dumps(review_notes, ensure_ascii=False), now(), inquiry_id))
        event(con, "ai_reply_created", "inquiry", inquiry_id,
              f"model={model}, intent={reply.intent}, knowledge_v={knowledge['version']}")
        return {"reply_status": "draft", "review_notes": review_notes, "intent": reply.intent, "model": model}


@app.put("/api/inquiries/{inquiry_id}/reply")
def edit_reply(inquiry_id: int, payload: ReplyEdit):
    with connect() as con:
        current = one(con, "inquiries", inquiry_id)
        if current["reply_status"] == "sent_manually":
            raise HTTPException(409, "已记录为手动发送的回复不可修改")
        if payload.reply_status == "sent_manually" and current["reply_status"] != "approved":
            raise HTTPException(409, "请先人工批准回复，再记录手动发送")
        if payload.reply_status == "sent_manually" and payload.reply_draft.strip() != current["reply_draft"]:
            raise HTTPException(409, "回复内容已变更；请重新保存并批准后再记录发送")
        if payload.reply_status in {"approved", "sent_manually"}:
            task = one(con, "publish_tasks", current["task_id"])
            draft = one(con, "drafts", task["draft_id"])
            product = one(con, "products", draft["product_id"])
            issues = reply_issues(product, current["question"], payload.reply_draft)
            if issues:
                raise HTTPException(422, {"issues": issues})
        con.execute("UPDATE inquiries SET reply_draft=?,reply_status=?,updated_at=? WHERE id=?",
                    (payload.reply_draft.strip(), payload.reply_status, now(), inquiry_id))
        event(con, "reply_status_changed", "inquiry", inquiry_id, payload.reply_status)
        return {"reply_status": payload.reply_status}


@app.put("/api/inquiries/{inquiry_id}/followup")
def followup(inquiry_id: int, payload: FollowupEdit):
    with connect() as con:
        one(con, "inquiries", inquiry_id)
        if payload.next_followup_at:
            try:
                datetime.fromisoformat(payload.next_followup_at)
            except ValueError as exc:
                raise HTTPException(422, "跟进时间必须是 ISO 日期时间") from exc
        con.execute("UPDATE inquiries SET stage=?,channel=?,next_followup_at=?,updated_at=? WHERE id=?",
                    (payload.stage, payload.channel, payload.next_followup_at, now(), inquiry_id))
        event(con, "followup_updated", "inquiry", inquiry_id, f"{payload.stage}, {payload.channel}")
        return {"stage": payload.stage}


@app.get("/api/export")
def export_data():
    return state()


@app.get("/api/browser/status")
def browser_status():
    return gateway.status()


@app.get("/api/browser/debug/snapshot")
async def browser_debug_snapshot():
    if os.getenv("WORKBENCH_DEBUG_BROWSER") != "1":
        raise HTTPException(404, "调试接口未启用")
    try:
        return {"snapshot": await gateway.debug_snapshot()}
    except Exception as exc:
        raise HTTPException(503, str(exc)[:200]) from exc


@app.get("/api/browser/debug/pages")
async def browser_debug_pages():
    if os.getenv("WORKBENCH_DEBUG_BROWSER") != "1":
        raise HTTPException(404, "调试接口未启用")
    try:
        return {"pages": await gateway.debug_pages()}
    except Exception as exc:
        raise HTTPException(503, str(exc)[:200]) from exc


@app.post("/api/browser/debug/disable")
def browser_debug_disable():
    if os.getenv("WORKBENCH_DEBUG_BROWSER") != "1":
        raise HTTPException(404, "调试接口未启用")
    os.environ.pop("WORKBENCH_DEBUG_BROWSER", None)
    return {"debug_enabled": False}


@app.post("/api/browser/debug/probe-inbox")
async def browser_debug_probe_inbox():
    if os.getenv("WORKBENCH_DEBUG_BROWSER") != "1":
        raise HTTPException(404, "调试接口未启用")
    try:
        return await gateway.probe_inbox()
    except Exception as exc:
        raise HTTPException(503, str(exc)[:200]) from exc


@app.post("/api/browser/connect")
async def browser_connect():
    return await gateway.start()


@app.post("/api/browser/disconnect")
async def browser_disconnect():
    return await gateway.disconnect()


@app.post("/api/browser/accounts/refresh")
async def browser_accounts_refresh():
    try:
        return await gateway.refresh_account()
    except Exception as exc:
        raise HTTPException(503, f"Chrome 账号读取失败：{str(exc)[:200]}") from exc


@app.get("/api/browser/conversations")
def browser_conversations():
    with connect() as con:
        return [dict(row) for row in con.execute(
            "SELECT * FROM browser_conversations ORDER BY needs_attention DESC,last_seen_at DESC,id DESC")]


@app.post("/api/browser/conversations/refresh")
async def browser_conversations_refresh():
    account_status = gateway.status()
    account_key = account_status.get("account_key")
    if account_status["state"] != "online" or not account_key or not account_status.get("display_name"):
        raise HTTPException(409, "请先连接并识别要同步会话的 Facebook 账号")
    try:
        items = await gateway.refresh_conversations()
    except Exception as exc:
        raise HTTPException(503, f"Marketplace 会话列表读取失败：{str(exc)[:200]}") from exc
    timestamp = now()
    visible_ids: list[int] = []
    with connect() as con:
        for item in items:
            previous = con.execute(
                "SELECT id,last_preview,needs_attention FROM browser_conversations "
                "WHERE account_key=? AND external_id=?",
                (account_key, item["external_id"])).fetchone()
            incoming = bool(item["preview"] and not item["preview"].startswith("You:"))
            if previous:
                needs_attention = bool(previous["needs_attention"]) or (
                    incoming and item["preview"] != previous["last_preview"])
                con.execute("""UPDATE browser_conversations
                    SET title=?,url=?,last_preview=?,needs_attention=?,last_seen_at=? WHERE id=?""",
                    (item["title"], item["url"], item["preview"], int(needs_attention), timestamp,
                     previous["id"]))
                visible_ids.append(previous["id"])
            else:
                inserted = con.execute("""INSERT INTO browser_conversations
                    (account_key,external_id,url,title,last_preview,needs_attention,first_seen_at,last_seen_at)
                    VALUES(?,?,?,?,?,?,?,?)""",
                    (account_key, item["external_id"], item["url"], item["title"],
                     item["preview"], int(incoming), timestamp, timestamp))
                visible_ids.append(inserted.lastrowid)
        event(con, "browser_conversation_sync", "browser_account", 0, f"visible={len(items)}")
        visible = ([dict(row) for row in con.execute(
            f"SELECT * FROM browser_conversations WHERE id IN ({','.join('?' for _ in visible_ids)}) "
            "ORDER BY needs_attention DESC,last_seen_at DESC,id DESC", visible_ids)]
            if visible_ids else [])
    return {"visible_count": len(items), "items": visible}


@app.post("/api/browser/conversations/{conversation_id}/read")
async def browser_conversation_read(conversation_id: int):
    with connect() as con:
        row = con.execute("SELECT * FROM browser_conversations WHERE id=?", (conversation_id,)).fetchone()
        if row is None:
            raise HTTPException(404, "会话不存在")
        account_status = gateway.status()
        if account_status["state"] != "online" or row["account_key"] != account_status.get("account_key"):
            raise HTTPException(409, "该会话所属账号当前未连接，请先切换到对应账号")
        url = row["url"]
    try:
        messages = await gateway.read_conversation(url)
    except Exception as exc:
        raise HTTPException(503, f"会话历史读取失败：{str(exc)[:200]}") from exc
    timestamp = now()
    occurrences: dict[str, int] = {}
    with connect() as con:
        for message in messages:
            basis = "\u241f".join((message["time_label"], message["sender"], message["text"]))
            occurrences[basis] = occurrences.get(basis, 0) + 1
            fingerprint = hashlib.sha256(f"{basis}\u241f{occurrences[basis]}".encode()).hexdigest()
            con.execute("""INSERT OR IGNORE INTO browser_messages
                (conversation_id,fingerprint,time_label,sender,direction,text,fetched_at)
                VALUES(?,?,?,?,?,?,?)""",
                (conversation_id, fingerprint, message["time_label"], message["sender"],
                 message["direction"], message["text"], timestamp))
        rows = [dict(item) for item in con.execute(
            "SELECT * FROM browser_messages WHERE conversation_id=? ORDER BY id", (conversation_id,))]
        event(con, "browser_history_sync", "browser_conversation", conversation_id,
              f"visible={len(messages)}")
    return {"messages": rows, "visible_count": len(messages)}


@app.post("/api/browser/conversations/{conversation_id}/handled")
def browser_conversation_handled(conversation_id: int):
    with connect() as con:
        row = con.execute("SELECT id FROM browser_conversations WHERE id=?", (conversation_id,)).fetchone()
        if row is None:
            raise HTTPException(404, "会话不存在")
        con.execute("UPDATE browser_conversations SET needs_attention=0 WHERE id=?", (conversation_id,))
        event(con, "browser_conversation_handled", "browser_conversation", conversation_id)
    return {"needs_attention": False}


@app.post("/api/browser/conversations/{conversation_id}/send")
async def browser_conversation_send(conversation_id: int, payload: BrowserSendRequest):
    body = payload.message.strip()
    with connect() as con:
        row = con.execute("SELECT * FROM browser_conversations WHERE id=?", (conversation_id,)).fetchone()
        if row is None:
            raise HTTPException(404, "会话不存在")
        account_status = gateway.status()
        if row["account_key"] != account_status.get("account_key") or account_status["state"] != "online":
            raise HTTPException(409, "目标账号不在线；请刷新账号状态")
        pending = con.execute("""SELECT id FROM browser_send_attempts
            WHERE conversation_id=? AND body=? AND state IN ('sending','uncertain')""",
            (conversation_id, body)).fetchone()
        if pending:
            raise HTTPException(409, "相同内容有未核实的发送记录，请先到 Facebook 核对，避免重复发送")
        cur = con.execute("""INSERT INTO browser_send_attempts
            (conversation_id,body,state,created_at,updated_at) VALUES(?,?,?,?,?)""",
            (conversation_id, body, "sending", now(), now()))
        attempt_id, url = cur.lastrowid, row["url"]
    try:
        outcome = await gateway.send_message(url, body)
        send_state = outcome["state"]
    except Exception:
        send_state = "uncertain"
    with connect() as con:
        con.execute("UPDATE browser_send_attempts SET state=?,updated_at=? WHERE id=?",
                    (send_state, now(), attempt_id))
        event(con, "browser_reply_attempt", "browser_conversation", conversation_id,
              f"attempt={attempt_id},state={send_state}")
        if send_state == "sent":
            con.execute("UPDATE browser_conversations SET needs_attention=0 WHERE id=?",
                        (conversation_id,))
    return {"attempt_id": attempt_id, "state": send_state,
            "message": "页面已确认发送" if send_state == "sent" else "发送结果不确定，请到 Facebook 核对；系统不会自动重试"}


@app.post("/api/browser/conversations/{conversation_id}/reply-ai")
def browser_conversation_reply_ai(conversation_id: int, payload: BrowserAIReplyRequest):
    if not payload.allow_external_processing:
        raise HTTPException(422, "请先确认消息已脱敏，并允许发送给配置的模型服务")
    with connect() as con:
        conversation = con.execute("SELECT id,account_key,title FROM browser_conversations WHERE id=?", (conversation_id,)).fetchone()
        if conversation is None:
            raise HTTPException(404, "会话不存在")
        account_status = gateway.status()
        if account_status["state"] != "online" or conversation["account_key"] != account_status.get("account_key"):
            raise HTTPException(409, "该会话所属账号当前未连接，请刷新账号状态")
        history = [dict(row) for row in con.execute("""SELECT sender,direction,time_label,text
            FROM browser_messages WHERE conversation_id=? ORDER BY id""", (conversation_id,))]
        if not history:
            raise HTTPException(409, "请先读取会话历史，再生成回复")
        if history[-1]["direction"] != "in":
            raise HTTPException(409, "最新一条消息不是买家发来的消息，请先核对会话")
        conversation_history = [
            {"sender": row["sender"], "direction": row["direction"],
             "time": row["time_label"], "text": row["text"]}
            for row in history
        ]
        try:
            context = reply_agent.plan_and_read(
                con, conversation, conversation_history, payload.product_id)
        except reply_agent.ProductMatchRequired as exc:
            raise HTTPException(409, str(exc)) from exc
        except ai.AIUnavailable as exc:
            raise HTTPException(503, str(exc)) from exc
        except OpenAIError as exc:
            raise HTTPException(503, "模型服务调用失败，请检查本机配置与网络") from exc
        product = context["product"]
        incoming = context["question"]
        knowledge = context["knowledge"]
        try:
            draft, model_id = ai.inquiry_reply(
                product if context["use_product_facts"] else None,
                incoming, conversation_history=context["conversation_history"],
                knowledge_context=knowledge)
        except ai.AIUnavailable as exc:
            raise HTTPException(503, str(exc)) from exc
        except OpenAIError as exc:
            raise HTTPException(503, "模型服务调用失败，请检查本机配置与网络") from exc
        answer = draft.answer.strip()
        if not answer or len(answer) > 2000:
            raise HTTPException(422, "AI 草稿为空或超过长度限制")
        issues = reply_issues(product, incoming, answer)
        if issues:
            raise HTTPException(422, {"issues": issues})
        event(con, "browser_ai_reply_drafted", "browser_conversation", conversation_id,
              f"product={product['id']},model={model_id},history={len(context['conversation_history'])},"
              f"knowledge_v={knowledge['version']},tools={','.join(context['tool_trace'])}")
    review_notes = list(dict.fromkeys([*draft.review_notes, *knowledge["missing_facts"]]))
    return {"answer": answer, "review_notes": review_notes, "model": model_id,
            "intent": draft.intent, "requires_review": bool(review_notes),
            "basis": {
                "product": f"{product['sku']} · {product['name']}",
                "history_count": len(context["conversation_history"]),
                "knowledge_version": knowledge["version"],
                "policies": [ai_knowledge.POLICY_FIELDS[key] for key in knowledge["policies"]],
                "reference_scripts": [item["title"] for item in knowledge["quick_replies"]],
                "tool_trace": context["tool_trace"],
            }}


_batch_task: asyncio.Task | None = None


async def _run_browser_publish_batch(task_ids: list[int]) -> None:
    for index, task_id in enumerate(task_ids):
        with connect() as con:
            task = one(con, "publish_tasks", task_id)
            draft = one(con, "drafts", task["draft_id"])
            product = one(con, "products", draft["product_id"])
            listing = {
                "sku": product["sku"], "title": draft["title"],
                "description": draft["description"],
                "price": f"{Decimal(product['price_minor']) / 100:.2f}",
                "category": "Jewellery and accessories",
                "condition": "New" if product["condition"].lower() == "new" else "Used – good",
                "image_path": product["image_path"],
            }
            con.execute("UPDATE publish_tasks SET status='preparing',updated_at=? WHERE id=?",
                        (now(), task_id))
            event(con, "browser_publish_started", "publish_task", task_id, task["account_alias"])
        try:
            outcome = await gateway.publish_listing(listing)
        except Exception as exc:
            outcome = {"state": "needs_attention", "note": f"网页操作停止：{str(exc)[:180]}"}
        with connect() as con:
            con.execute("""UPDATE publish_tasks
                SET status=?,marketplace_url=?,note=?,updated_at=? WHERE id=?""",
                (outcome["state"], outcome.get("url"), outcome["note"], now(), task_id))
            event(con, "browser_publish_result", "publish_task", task_id, outcome["state"])
        if outcome["state"] != "posted":
            with connect() as con:
                for remaining in task_ids[index + 1:]:
                    con.execute("UPDATE publish_tasks SET status='needs_attention',note=?,updated_at=? WHERE id=?",
                                ("前一任务未确认发布，队列已暂停", now(), remaining))
                    event(con, "browser_publish_paused", "publish_task", remaining)
            break


@app.post("/api/browser/publish/start")
async def browser_publish_start(payload: PublishBatchRequest):
    global _batch_task
    if not payload.confirm:
        raise HTTPException(422, "请在工作台确认商品、图片与目标账号后再开始")
    if payload.account_keys != ["chrome-default"]:
        raise HTTPException(422, "当前版本只支持已连接 Chrome 默认资料目录的一个账号")
    if gateway.status()["state"] != "online":
        raise HTTPException(409, "当前 Chrome 账号不在线")
    if _batch_task is not None and not _batch_task.done():
        raise HTTPException(409, "已有网页上架任务正在执行")
    product_ids = list(dict.fromkeys(payload.product_ids))
    task_ids = []
    with connect() as con:
        products = []
        for product_id in product_ids:
            product = one(con, "products", product_id)
            if not product["active"]:
                raise HTTPException(409, f"{product['sku']} 已从当前商品表移除，请重新导入后再发布")
            if product["currency"] != "USD" or product["stock"] < 1:
                raise HTTPException(422, f"{product['sku']} 的币种或库存不符合当前上架条件")
            if not product["listing_title"] or not product["listing_description"]:
                raise HTTPException(422, f"{product['sku']} 缺少商品表中的 listing_title / listing_description")
            if not product["image_path"] or not Path(product["image_path"]).is_file():
                raise HTTPException(422, f"{product['sku']} 尚未上传本地图片")
            issues = draft_issues(product, product["listing_title"], product["listing_description"])
            if issues:
                raise HTTPException(422, {"sku": product["sku"], "issues": issues})
            existing = con.execute("""SELECT t.id FROM publish_tasks t
                JOIN drafts d ON d.id=t.draft_id WHERE d.product_id=? AND t.account_alias=?
                AND t.status IN ('queued','preparing','posted','uncertain') LIMIT 1""",
                (product_id, "chrome-default")).fetchone()
            if existing:
                raise HTTPException(409, f"{product['sku']} 在此账号已有待处理、已发布或待核实任务")
            products.append(product)
        for product in products:
            version = con.execute("SELECT COALESCE(MAX(version),0)+1 FROM drafts WHERE product_id=?",
                                  (product["id"],)).fetchone()[0]
            cur = con.execute("""INSERT INTO drafts
                (product_id,title,description,language,status,version,created_at,approved_at,generation_source)
                VALUES(?,?,?,?,?,?,?,?,?)""",
                (product["id"], product["listing_title"], product["listing_description"], "en",
                 "approved", version, now(), now(), "imported_sheet"))
            task = con.execute("""INSERT INTO publish_tasks
                (draft_id,account_alias,status,created_at,updated_at) VALUES(?,?,?,?,?)""",
                (cur.lastrowid, "chrome-default", "queued", now(), now()))
            task_ids.append(task.lastrowid)
            event(con, "browser_batch_queued", "publish_task", task.lastrowid,
                  f"sku={product['sku']},account=chrome-default")
    _batch_task = asyncio.create_task(_run_browser_publish_batch(task_ids))
    return {"task_ids": task_ids, "state": "queued"}


@app.post("/api/browser/publish/{task_id}/verify")
async def browser_publish_verify(task_id: int):
    with connect() as con:
        task = one(con, "publish_tasks", task_id)
        if task["status"] != "uncertain":
            raise HTTPException(409, "只有待核实的任务可以重新核验")
        draft = one(con, "drafts", task["draft_id"])
        title = draft["title"]
    try:
        outcome = await gateway.verify_listing(title)
    except Exception as exc:
        raise HTTPException(503, str(exc)[:200]) from exc
    with connect() as con:
        con.execute("""UPDATE publish_tasks SET status=?,marketplace_url=?,note=?,updated_at=? WHERE id=?""",
                    (outcome["state"], outcome.get("url"), outcome["note"], now(), task_id))
        event(con, "browser_publish_verified", "publish_task", task_id, outcome["state"])
    return outcome
